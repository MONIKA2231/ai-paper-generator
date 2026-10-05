import json
import re
import io
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session
from pypdf import PdfReader
from docx import Document

from ..db import get_db
from ..models import (
    Subject,
    Syllabus,
    Question,
    Blueprint,
    Paper,
    AnswerKey,
    PreviousPaper,
    VaultItem,
    User,
    Setting,
    AuditLog,
    LoginSession,
)
from datetime import datetime
from ..security import current_user, verify_password, hash_password
from ..services.ai import (
    analyze_syllabus,
    classify_question,
    generate_answer,
    assistant_answer,
    generate_paper as ai_generate_paper,
)


router = APIRouter(tags=["AIQPG Modules"])


# =========================================================
# INPUT SCHEMAS
# =========================================================

class SyllabusIn(BaseModel):
    subject_id: int
    title: str
    text: str = ""
    file_name: str = ""


class QuestionIn(BaseModel):
    subject_id: int
    syllabus_id: int | None = None
    question_text: str
    unit: str = ""
    topic: str = ""
    marks: int = 1


class BlueprintIn(BaseModel):
    subject_id: int
    syllabus_id: int | None = None
    blueprint_name: str
    exam_type: str = "Internal Assessment"
    duration: str = "90 Minutes"
    total_marks: int = 50
    total_questions: int = 25

    question_pattern: dict = {}
    unit_weightage: dict = {}
    difficulty_distribution: dict = {}
    bloom_distribution: dict = {}
    co_distribution: dict = {}


class PaperGenIn(BaseModel):
    subject_id: int
    syllabus_id: int | None = None
    blueprint_id: int | None = None
    title: str = "AI Generated Question Paper"
    exam_type: str = "Internal Assessment"
    duration: str = "90 Minutes"
    total_marks: int = 50
    total_questions: int = 25
    section_a_questions: int = 0
    section_a_marks: int = 2
    section_b_questions: int = 0
    section_b_marks: int = 8
    choice_mode: str = "No Choice"
    generation_mode: str = "hybrid"
    topics: list[str] = []
    instructions: str = ""


class AnswerGenIn(BaseModel):
    paper_id: int | None = None
    question_id: int | None = None
    question_text: str
    marks: int = 1


# =========================================================
# SUBJECT MANAGEMENT
# =========================================================

@router.get("/subjects")
def list_subjects(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    subjects = db.query(Subject).order_by(Subject.id.desc()).all()
    return [
        {
            "id": s.id,
            "name": s.name,
            "code": s.code,
            "department": s.department,
            "semester": s.semester,
            "academic_year": s.academic_year,
            "credits": s.credits,
            "description": s.description,
        }
        for s in subjects
    ]


class SubjectIn(BaseModel):
    name: str
    code: str
    department: str = ""
    semester: str = ""
    academic_year: str = ""
    credits: int = 0
    description: str = ""


@router.post("/subjects")
def create_subject(
    data: SubjectIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    name = data.name.strip()
    code = data.code.strip().upper()
    if not name or not code:
        raise HTTPException(400, "Subject name and code are required.")

    if db.query(Subject).filter(func.lower(Subject.name) == name.lower()).first():
        raise HTTPException(400, "Subject already exists.")
    if db.query(Subject).filter(func.lower(Subject.code) == code.lower()).first():
        raise HTTPException(400, "Subject code already exists.")

    subject = Subject(
        name=name,
        code=code,
        department=data.department.strip(),
        semester=data.semester.strip(),
        academic_year=data.academic_year.strip(),
        credits=data.credits,
        description=data.description.strip(),
        created_by=user.id,
    )
    db.add(subject)
    db.commit()
    db.refresh(subject)
    return {"message": "Subject created", "id": subject.id}


@router.put("/subjects/{subject_id}")
def update_subject(
    subject_id: int,
    data: SubjectIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    subject = db.get(Subject, subject_id)
    if not subject:
        raise HTTPException(404, "Subject not found")

    subject.name = data.name.strip()
    subject.code = data.code.strip().upper()
    subject.department = data.department.strip()
    subject.semester = data.semester.strip()
    subject.academic_year = data.academic_year.strip()
    subject.credits = data.credits
    subject.description = data.description.strip()
    db.commit()
    return {"message": "Subject updated"}


@router.delete("/subjects/{subject_id}")
def delete_subject(
    subject_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    subject = db.get(Subject, subject_id)
    if not subject:
        raise HTTPException(404, "Subject not found")
    db.delete(subject)
    db.commit()
    return {"message": "Subject deleted"}


# =========================================================
# DASHBOARD
# =========================================================

@router.get("/dashboard")
def dashboard(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    subject_count = db.query(Subject).count()
    question_count = db.query(Question).count()
    paper_count = db.query(Paper).count()
    syllabus_count = db.query(Syllabus).count()

    coverage = 0
    if subject_count:
        coverage = round(min(100, (syllabus_count / subject_count) * 100))

    recent_papers = db.query(Paper).order_by(Paper.id.desc()).limit(5).all()
    recent_papers_list = [
        {
            "id": p.id,
            "title": p.title,
            "exam_type": p.exam_type,
            "total_marks": p.total_marks,
            "total_questions": p.total_questions,
            "status": p.status,
            "created_at": p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else "",
        }
        for p in recent_papers
    ]

    recent_activity = db.query(AuditLog).order_by(AuditLog.id.desc()).limit(5).all()
    recent_activity_list = [
        {
            "id": a.id,
            "action": a.action,
            "details": a.details,
            "created_at": a.created_at.strftime("%Y-%m-%d %H:%M") if a.created_at else "",
        }
        for a in recent_activity
    ]

    easy_q = db.query(Question).filter(Question.difficulty == "Easy").count()
    medium_q = db.query(Question).filter(Question.difficulty == "Medium").count()
    hard_q = db.query(Question).filter(Question.difficulty == "Hard").count()
    sec_a_q = db.query(Question).filter(Question.marks == 2).count()
    sec_b_q = db.query(Question).filter(Question.marks == 8).count()

    return {
        "subjects": subject_count,
        "syllabi": syllabus_count,
        "questions": question_count,
        "papers": paper_count,
        "coverage": coverage,
        "recent_papers": recent_papers_list,
        "recent_activity": recent_activity_list,
        "question_stats": {
            "easy": easy_q,
            "medium": medium_q,
            "hard": hard_q,
            "sec_a": sec_a_q,
            "sec_b": sec_b_q,
        },
    }


# =========================================================
# SETTINGS & PASSWORD
# =========================================================

@router.get("/settings")
def get_user_settings(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    rows = db.query(Setting).filter(Setting.user_id == user.id).all()
    res = {}
    for r in rows:
        try:
            res[r.key] = json.loads(r.value)
        except Exception:
            res[r.key] = r.value
    return res

@router.post("/settings")
def save_user_settings(
    payload: dict,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    settings_dict = payload.get("settings") if isinstance(payload.get("settings"), dict) else payload
    for key, val in settings_dict.items():
        val_str = json.dumps(val) if isinstance(val, (dict, list, bool)) else str(val)
        setting = db.query(Setting).filter(Setting.user_id == user.id, Setting.key == key).first()
        if setting:
            setting.value = val_str
            setting.updated_at = datetime.utcnow()
        else:
            db.add(Setting(user_id=user.id, key=key, value=val_str))
    db.commit()
    return {"message": "Settings saved successfully."}

class DirectChangePasswordIn(BaseModel):
    current_password: str
    new_password: str

@router.post("/change-password")
def change_password_direct(
    data: DirectChangePasswordIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    current_password = str(data.current_password or '')
    new_password = str(data.new_password or '')

    if not current_password:
        raise HTTPException(status_code=400, detail="Current password is required.")
    if len(new_password) < 8:
        raise HTTPException(status_code=400, detail="New password must contain at least 8 characters.")
    if not verify_password(current_password, user.password_hash):
        raise HTTPException(status_code=400, detail="Current password is incorrect.")

    user.password_hash = hash_password(new_password)
    db.add(user)
    db.add(AuditLog(user_id=user.id, action="CHANGE_PASSWORD", details="User changed account password"))
    db.commit()

    return {"message": "Password changed successfully."}


# =========================================================
# ADMIN CONSOLE
# =========================================================

def require_admin(user):
    if getattr(user, "role", "") != "admin":
        raise HTTPException(403, "Administrator access required.")
    return user


@router.get("/admin/users")
def admin_users(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    require_admin(user)
    users = db.query(User).order_by(User.id.desc()).all()
    return [
        {
            "id": u.id,
            "name": u.name,
            "email": u.email,
            "role": u.role,
            "department": u.department,
            "last_login": u.last_login,
            "last_active": u.last_active,
        }
        for u in users
    ]


@router.get("/admin/logs")
def admin_logs(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    require_admin(user)
    logs = db.query(AuditLog).order_by(AuditLog.id.desc()).limit(200).all()
    return [
        {
            "id": x.id,
            "user_id": x.user_id,
            "action": x.action,
            "details": x.details,
            "created_at": x.created_at,
        }
        for x in logs
    ]


@router.get("/admin/sessions")
def admin_sessions(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    require_admin(user)
    sessions = db.query(LoginSession).order_by(LoginSession.id.desc()).limit(200).all()
    return [
        {
            "id": x.id,
            "user_id": x.user_id,
            "login_at": x.login_at,
            "logout_at": x.logout_at,
            "status": x.status,
            "ip_address": x.ip_address,
        }
        for x in sessions
    ]


@router.get("/admin/tables")
def admin_tables(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    require_admin(user)
    return {
        "tables": [
            "answer_keys",
            "audit_logs",
            "blueprints",
            "login_sessions",
            "papers",
            "previous_papers",
            "questions",
            "settings",
            "subjects",
            "syllabi",
            "users",
            "vault_items",
        ]
    }


# =========================================================
# SYLLABUS
# =========================================================

@router.post("/syllabi/upload")
async def upload_syllabus(
    subject_id: int,
    title: str = "",
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    raw = await file.read()
    name = (file.filename or "").lower()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(
            (p.extract_text() or "")
            for p in reader.pages
        )

    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(raw))
        text = "\n".join(
            p.text
            for p in doc.paragraphs
        )

    elif name.endswith(".txt"):
        text = raw.decode(
            "utf-8",
            errors="ignore"
        )

    elif name.endswith(".doc"):
        import shutil
        import subprocess
        import tempfile
        import os

        soffice = shutil.which("soffice")

        if not soffice:
            raise HTTPException(
                400,
                "DOC files require LibreOffice (soffice) on the server. "
                "Use DOCX or PDF if LibreOffice is unavailable.",
            )

        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(
                td,
                file.filename
            )

            with open(src, "wb") as f:
                f.write(raw)

            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    td,
                    src,
                ],
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )

            converted = os.path.join(
                td,
                Path(file.filename).stem + ".docx"
            )

            if not os.path.exists(converted):
                raise HTTPException(
                    400,
                    "Unable to read the DOC file."
                )

            doc = Document(converted)

            text = "\n".join(
                p.text
                for p in doc.paragraphs
            )

    else:
        raise HTTPException(
            400,
            "Only PDF, DOC, DOCX and TXT files are supported.",
        )

    if not text.strip():
        raise HTTPException(
            400,
            "No readable text was found in the syllabus file.",
        )

    analysis = analyze_syllabus(text)

    syllabus = Syllabus(
        subject_id=subject_id,
        title=(
            title.strip()
            or file.filename
        ),
        file_name=file.filename,
        extracted_text=text,
        units=json.dumps(
            analysis["units"]
        ),
        learning_outcomes=json.dumps(
            analysis["learning_outcomes"]
        ),
        total_units=analysis["total_units"],
        total_topics=analysis["total_topics"],
        status="analyzed",
    )

    db.add(syllabus)
    db.commit()
    db.refresh(syllabus)

    return {
        "id": syllabus.id,
        "analysis": analysis,
    }


@router.post("/syllabi/analyze")
def analyze_syllabus_text(
    data: SyllabusIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    analysis = analyze_syllabus(data.text)

    syllabus = Syllabus(
        subject_id=data.subject_id,
        title=data.title.strip() if data.title else "Syllabus",
        file_name=data.file_name or "Manual Entry",
        extracted_text=data.text,
        units=json.dumps(analysis["units"]),
        learning_outcomes=json.dumps(analysis["learning_outcomes"]),
        total_units=analysis["total_units"],
        total_topics=analysis["total_topics"],
        status="analyzed",
    )
    db.add(syllabus)
    db.commit()
    db.refresh(syllabus)

    return {
        "id": syllabus.id,
        "analysis": analysis,
        "message": "Syllabus analyzed and saved successfully.",
    }


@router.get("/syllabi")
def list_syllabi(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    syllabi = db.query(Syllabus).all()

    return [
        {
            "id": s.id,
            "subject_id": s.subject_id,
            "title": s.title,
            "total_units": s.total_units,
            "total_topics": s.total_topics,
            "status": s.status,
        }
        for s in syllabi
    ]


@router.get("/syllabi/{sid}")
def get_syllabus(
    sid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    s = db.get(Syllabus, sid)

    if not s:
        raise HTTPException(
            404,
            "Syllabus not found"
        )

    return {
        "id": s.id,
        "subject_id": s.subject_id,
        "title": s.title,
        "units": json.loads(
            s.units or "[]"
        ),
        "learning_outcomes": json.loads(
            s.learning_outcomes or "[]"
        ),
        "total_units": s.total_units,
        "total_topics": s.total_topics,
        "status": s.status,
    }


@router.put("/syllabi/{sid}")
def update_syllabus(
    sid: int,
    data: SyllabusIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    s = db.get(Syllabus, sid)

    if not s:
        raise HTTPException(
            404,
            "Syllabus not found"
        )

    analysis = analyze_syllabus(data.text)

    s.subject_id = data.subject_id
    s.title = data.title
    s.file_name = data.file_name
    s.extracted_text = data.text
    s.units = json.dumps(
        analysis["units"]
    )
    s.learning_outcomes = json.dumps(
        analysis["learning_outcomes"]
    )
    s.total_units = analysis["total_units"]
    s.total_topics = analysis["total_topics"]
    s.status = "analyzed"

    db.commit()
    db.refresh(s)

    return {
        "id": s.id,
        "analysis": analysis,
    }


@router.delete("/syllabi/{sid}")
def delete_syllabus(
    sid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    s = db.get(Syllabus, sid)

    if not s:
        raise HTTPException(
            404,
            "Syllabus not found"
        )

    db.delete(s)
    db.commit()

    return {
        "message": "Syllabus deleted"
    }


# =========================================================
# QUESTION BANK
# =========================================================

@router.post("/questions")
def create_question(
    data: QuestionIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    meta = classify_question(
        data.question_text
    )

    payload = data.model_dump()

    allowed_meta = {
        "syllabus_id", "unit", "topic", "question_type",
        "difficulty", "marks", "bloom_level", "course_outcome",
        "keywords", "options", "answer",
    }
    payload.update({k: v for k, v in (meta or {}).items() if k in allowed_meta})
    payload["created_by"] = user.id

    question = Question(**payload)

    db.add(question)
    db.commit()
    db.refresh(question)

    return question.__dict__

@router.get("/questions")
def list_questions(
    subject_name: str | None = None,
    subject_id: int | None = None,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    # Filter by subject ID if provided
    if subject_id:
        query = db.query(Question).filter(
            Question.subject_id == subject_id
        )

    # Otherwise filter by subject name
    elif subject_name and subject_name.strip():

        subject = (
            db.query(Subject)
            .filter(
                func.lower(func.trim(Subject.name))
                == subject_name.strip().lower()
            )
            .first()
        )

        # Subject does not exist
        if not subject:
            return []

        query = db.query(Question).filter(
            Question.subject_id == subject.id
        )

    # Never show questions from all subjects
    else:
        return []

    questions = (
        query
        .order_by(Question.id.desc())
        .all()
    )

    return [
        {
            "id": q.id,
            "subject_id": q.subject_id,
            "unit": q.unit,
            "topic": q.topic,
            "question_text": q.question_text,
            "question_type": q.question_type,
            "difficulty": q.difficulty,
            "marks": q.marks,
            "bloom_level": q.bloom_level,
            "course_outcome": q.course_outcome,
            "duplicate_status": q.duplicate_status,
            "mapping_status": q.mapping_status,
            "status": q.status,
        }
        for q in questions
    ]


@router.post("/questions/upload")
async def upload_questions(
    subject_name: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    """Upload and analyze a question bank for one subject."""

    subject_name = subject_name.strip()
    if not subject_name:
        raise HTTPException(400, "Subject name is required.")

    raw = await file.read()
    if not raw:
        raise HTTPException(400, "Uploaded file is empty.")

    filename = (file.filename or "").lower()

    # Find the subject. Create it when it does not already exist.
    subject = (
        db.query(Subject)
        .filter(
            func.lower(func.trim(Subject.name)) == subject_name.lower()
        )
        .first()
    )

    if not subject:
        code = re.sub(r"[^A-Za-z0-9]+", "", subject_name).upper()[:20] or "SUBJECT"
        original_code = code
        counter = 1
        while db.query(Subject).filter(Subject.code == code).first():
            code = f"{original_code}{counter}"
            counter += 1

        subject = Subject(
            name=subject_name,
            code=code,
            department="",
            semester="",
            academic_year="",
            credits=0,
            description="Created from Question Bank upload",
            created_by=user.id,
        )
        db.add(subject)
        db.commit()
        db.refresh(subject)

    # Extract text from the supported question-bank formats.
    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in doc.paragraphs)

    elif filename.endswith(".txt"):
        text = raw.decode("utf-8", errors="ignore")

    elif filename.endswith(".csv"):
        import csv
        decoded = raw.decode("utf-8", errors="ignore")
        rows = csv.reader(io.StringIO(decoded))
        text = "\n".join(
            " ".join(cell.strip() for cell in row if cell and cell.strip())
            for row in rows
        )

    elif filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(400, "XLSX processing requires openpyxl. Install it with: pip install openpyxl")

        try:
            workbook = load_workbook(
                io.BytesIO(raw),
                read_only=True,
                data_only=True,
            )
            rows = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value).strip()
                        for value in row
                        if value is not None and str(value).strip()
                    ]
                    if values:
                        rows.append(" ".join(values))
            text = "\n".join(rows)
        except Exception as exc:
            raise HTTPException(400, f"Unable to read XLSX file: {exc}")

    else:
        raise HTTPException(
            400,
            "Unsupported file format. Use PDF, DOCX, XLSX, CSV or TXT.",
        )

    if not text.strip():
        raise HTTPException(400, "No readable questions were found in the file.")

    # Existing questions for this subject are used for duplicate checking.
    existing_questions = (
        db.query(Question)
        .filter(Question.subject_id == subject.id)
        .all()
    )

    created_count = 0
    duplicate_count = 0
    mapped_count = 0
    review_count = 0

    if filename.endswith((".pdf", ".docx", ".txt")):
        # Remove section headers entirely from the text
        text = re.sub(r"(?im)^\s*(?:section|part|module|unit)\s*[-_:]?\s*[a-z0-9\s,]+(?:\s*\(.*?\))?\s*$", "", text)
        
        # Split by question numbers at the start of a line
        blocks = re.split(r"(?im)^\s*(?:Q(?:uestion)?\s*)?\d+[\.\)\:\-]\s+", text)
        
        # If it found question numbers, use blocks (ignoring the first chunk which is usually intro text)
        if len(blocks) > 2:
            lines = []
            for b in blocks[1:]:
                # Flatten the block into a single line, replacing newlines with spaces
                b_flat = re.sub(r"\s+", " ", b).strip()
                if b_flat:
                    lines.append(b_flat)
        else:
            # Fallback to line-by-line if no clear numbering exists
            lines = [line.strip() for line in text.splitlines() if len(line.strip()) > 8]
    else:
        # For CSV/XLSX, each row is a line
        lines = [line.strip() for line in text.splitlines() if len(line.strip()) > 8]

    for line in lines:
        # Remove common numbering such as 1., Q1., Question 1), etc. in case it wasn't stripped
        line = re.sub(
            r"^\s*(?:Q(?:uestion)?\s*)?\d+[\.\):\-]\s*",
            "",
            line,
            flags=re.I,
        ).strip()
        
        if len(line) < 8:
            continue

        # AI classification.
        raw_meta = classify_question(line) or {}

        # Keep only fields that actually belong to the Question model.
        allowed_meta = {
            "syllabus_id",
            "unit",
            "topic",
            "question_type",
            "difficulty",
            "marks",
            "bloom_level",
            "course_outcome",
            "keywords",
            "options",
            "answer",
        }
        meta = {key: value for key, value in raw_meta.items() if key in allowed_meta}

        # Duplicate detection within the uploaded subject.
        new_words = set(re.findall(r"\w+", line.lower()))
        duplicate_status = "unique"

        for existing in existing_questions:
            old_words = set(
                re.findall(r"\w+", (existing.question_text or "").lower())
            )
            similarity = len(new_words & old_words) / max(1, len(new_words | old_words))
            if similarity >= 0.65:
                duplicate_status = "duplicate"
                break

        if duplicate_status == "duplicate":
            duplicate_count += 1

        # Syllabus mapping status.
        mapping_status = "mapped" if (meta.get("unit") or meta.get("topic")) else "needs_review"
        if mapping_status == "mapped":
            mapped_count += 1
        else:
            review_count += 1

        question = Question(
            subject_id=subject.id,
            question_text=line,
            syllabus_id=meta.get("syllabus_id"),
            unit=meta.get("unit", ""),
            topic=meta.get("topic", ""),
            question_type=meta.get("question_type", "Descriptive"),
            difficulty=meta.get("difficulty", "Medium"),
            marks=meta.get("marks", 1),
            bloom_level=meta.get("bloom_level", "Understand"),
            course_outcome=meta.get("course_outcome", "CO1"),
            keywords=meta.get("keywords", ""),
            options=meta.get("options", "[]"),
            answer=meta.get("answer", ""),
            source_file=file.filename or "",
            duplicate_status=duplicate_status,
            mapping_status=mapping_status,
            status="approved",
            created_by=user.id,
        )

        db.add(question)
        existing_questions.append(question)
        created_count += 1

    db.commit()

    return {
        "message": f"{created_count} questions processed for {subject.name}.",
        "count": created_count,
        "subject_id": subject.id,
        "subject_name": subject.name,
        "duplicates": duplicate_count,
        "mapped": mapped_count,
        "needs_review": review_count,
    }


@router.delete("/questions/{qid}")
def delete_question(
    qid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    question = db.get(
        Question,
        qid
    )

    if not question:
        raise HTTPException(
            404,
            "Question not found"
        )

    db.delete(question)
    db.commit()

    return {"message": "Question deleted"}

@router.delete("/questions/subject/{subject_id}/clear")
def clear_subject_questions(
    subject_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    subject = db.get(Subject, subject_id)
    if not subject:
        raise HTTPException(404, "Subject not found")

    deleted = db.query(Question).filter(Question.subject_id == subject_id).delete()
    db.commit()

    return {"message": f"Successfully cleared {deleted} questions for {subject.name}."}


@router.post("/questions/check-duplicates")
def duplicates(
    subject_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    questions = (
        db.query(Question)
        .filter(
            Question.subject_id == subject_id
        )
        .all()
    )

    groups = []

    for i, a in enumerate(questions):

        for b in questions[i + 1:]:

            words_a = set(
                re.findall(
                    r"\w+",
                    a.question_text.lower()
                )
            )

            words_b = set(
                re.findall(
                    r"\w+",
                    b.question_text.lower()
                )
            )

            score = (
                len(words_a & words_b)
                / max(
                    1,
                    len(words_a | words_b)
                )
            )

            if score >= 0.65:
                groups.append(
                    {
                        "question_a": a.id,
                        "question_b": b.id,
                        "similarity": round(
                            score * 100,
                            1
                        ),
                    }
                )

    return {
        "duplicates": groups,
        "count": len(groups),
    }


@router.post("/questions/map-syllabus")
def map_syllabus(
    subject_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    questions = (
        db.query(Question)
        .filter(
            Question.subject_id == subject_id
        )
        .all()
    )

    for question in questions:

        if question.unit or question.topic:
            question.mapping_status = "mapped"
        else:
            question.mapping_status = "needs_review"

    db.commit()

    return {
        "mapped": sum(
            q.mapping_status == "mapped"
            for q in questions
        ),
        "needs_review": sum(
            q.mapping_status != "mapped"
            for q in questions
        ),
    }


@router.post("/questions/{qid}/approve")
def approve_question(
    qid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    question = db.get(
        Question,
        qid
    )

    if not question:
        raise HTTPException(
            404,
            "Question not found"
        )

    question.status = "approved"

    db.commit()

    return {
        "message": "Question approved"
    }


@router.post("/questions/{qid}/reject")
def reject_question(
    qid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    question = db.get(
        Question,
        qid
    )

    if not question:
        raise HTTPException(
            404,
            "Question not found"
        )

    question.status = "rejected"

    db.commit()

    return {
        "message": "Question rejected"
    }


# =========================================================
# BLUEPRINT DESIGNER
# =========================================================

@router.post("/blueprints")
def create_blueprint(
    data: BlueprintIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    # IMPORTANT:
    # data.model_dump() already contains the dictionary fields.
    # Convert those fields to JSON inside the payload instead
    # of passing them twice to Blueprint().

    payload = data.model_dump()

    payload["question_pattern"] = json.dumps(
        data.question_pattern
    )

    payload["unit_weightage"] = json.dumps(
        data.unit_weightage
    )

    payload["difficulty_distribution"] = json.dumps(
        data.difficulty_distribution
    )

    payload["bloom_distribution"] = json.dumps(
        data.bloom_distribution
    )

    payload["co_distribution"] = json.dumps(
        data.co_distribution
    )

    payload["created_by"] = user.id

    blueprint = Blueprint(
        **payload
    )

    db.add(blueprint)
    db.commit()
    db.refresh(blueprint)

    return {
        "id": blueprint.id,
        "name": blueprint.blueprint_name,
        "status": blueprint.status,
    }


@router.get("/blueprints")
def list_blueprints(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    blueprints = (
        db.query(Blueprint)
        .all()
    )

    return [
        {
            "id": b.id,
            "subject_id": b.subject_id,
            "blueprint_name": b.blueprint_name,
            "exam_type": b.exam_type,
            "total_marks": b.total_marks,
            "total_questions": b.total_questions,
            "status": b.status,
        }
        for b in blueprints
    ]


@router.get("/blueprints/{bid}")
def get_blueprint(
    bid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    blueprint = db.get(
        Blueprint,
        bid
    )

    if not blueprint:
        raise HTTPException(
            404,
            "Blueprint not found"
        )

    return {
        "id": blueprint.id,
        "subject_id": blueprint.subject_id,
        "syllabus_id": blueprint.syllabus_id,
        "blueprint_name": blueprint.blueprint_name,
        "exam_type": blueprint.exam_type,
        "duration": blueprint.duration,
        "total_marks": blueprint.total_marks,
        "total_questions": blueprint.total_questions,

        "question_pattern": json.loads(
            blueprint.question_pattern or "{}"
        ),

        "unit_weightage": json.loads(
            blueprint.unit_weightage or "{}"
        ),

        "difficulty_distribution": json.loads(
            blueprint.difficulty_distribution or "{}"
        ),

        "bloom_distribution": json.loads(
            blueprint.bloom_distribution or "{}"
        ),

        "co_distribution": json.loads(
            blueprint.co_distribution or "{}"
        ),

        "status": blueprint.status,
    }


@router.put("/blueprints/{bid}")
def update_blueprint(
    bid: int,
    data: BlueprintIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    blueprint = db.get(
        Blueprint,
        bid
    )

    if not blueprint:
        raise HTTPException(
            404,
            "Blueprint not found"
        )

    payload = data.model_dump()

    payload["question_pattern"] = json.dumps(
        data.question_pattern
    )

    payload["unit_weightage"] = json.dumps(
        data.unit_weightage
    )

    payload["difficulty_distribution"] = json.dumps(
        data.difficulty_distribution
    )

    payload["bloom_distribution"] = json.dumps(
        data.bloom_distribution
    )

    payload["co_distribution"] = json.dumps(
        data.co_distribution
    )

    for key, value in payload.items():
        setattr(
            blueprint,
            key,
            value
        )

    db.commit()
    db.refresh(blueprint)

    return {
        "id": blueprint.id,
        "message": "Blueprint updated",
    }


@router.delete("/blueprints/{bid}")
def delete_blueprint(
    bid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    blueprint = db.get(
        Blueprint,
        bid
    )

    if not blueprint:
        raise HTTPException(
            404,
            "Blueprint not found"
        )

    db.delete(blueprint)
    db.commit()

    return {
        "message": "Blueprint deleted"
    }


@router.post("/blueprints/validate")
def validate_blueprint(
    data: BlueprintIn
):
    valid = (
        data.total_marks > 0
        and data.total_questions > 0
    )

    return {
        "valid": valid,
        "checks": {
            "total_marks": data.total_marks > 0,
            "question_count": data.total_questions > 0,
            "unit_weightage": True,
            "difficulty_balance": True,
            "bloom_balance": True,
            "co_coverage": True,
        },
    }


@router.post("/blueprints/generate")
def generate_blueprint(
    data: BlueprintIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    return create_blueprint(
        data,
        db,
        user
    )


@router.get("/blueprints/{bid}/preview")
def preview_blueprint(
    bid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    return get_blueprint(
        bid,
        db,
        user
    )


# =========================================================
# PAPER GENERATION
# =========================================================

@router.post("/papers/generate")
def generate_paper(
    data: PaperGenIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    subject = db.get(Subject, data.subject_id)
    if not subject:
        raise HTTPException(404, "Subject not found")

    questions = (
        db.query(Question)
        .filter(
            Question.subject_id == data.subject_id,
            Question.status == "approved",
        )
        .limit(200)
        .all()
    )

    if data.generation_mode == "question_bank":
        content = []
        for i, question in enumerate(questions[:data.total_questions], 1):
            content.append({
                "number": i,
                "section": "Section A" if question.marks <= 2 else "Section B",
                "question": question.question_text,
                "marks": question.marks,
                "unit": question.unit,
                "topic": question.topic,
                "difficulty": question.difficulty,
                "bloom_level": question.bloom_level,
                "co": question.course_outcome,
            })
    else:
        instructions = data.instructions or f"Section A: {data.section_a_questions} questions of {data.section_a_marks} marks each. Section B: {data.section_b_questions} questions of {data.section_b_marks} marks each."
        
        syllabus_text = ""
        if data.syllabus_id:
            syllabus = db.get(Syllabus, data.syllabus_id)
            if syllabus:
                syllabus_text = syllabus.extracted_text

        blueprint_text = ""
        if data.blueprint_id:
            blueprint = db.get(Blueprint, data.blueprint_id)
            if blueprint:
                blueprint_text = json.dumps({
                    "question_pattern": json.loads(blueprint.question_pattern or "{}"),
                    "unit_weightage": json.loads(blueprint.unit_weightage or "{}"),
                    "difficulty_distribution": json.loads(blueprint.difficulty_distribution or "{}"),
                    "bloom_distribution": json.loads(blueprint.bloom_distribution or "{}"),
                    "co_distribution": json.loads(blueprint.co_distribution or "{}")
                })

        content = ai_generate_paper(
            subject=subject,
            questions=questions if data.generation_mode == "hybrid" else [],
            total_marks=data.total_marks,
            total_questions=data.total_questions,
            duration=data.duration,
            choice_mode=data.choice_mode,
            topics=data.topics,
            instructions=instructions,
            syllabus_text=syllabus_text,
            blueprint_text=blueprint_text
        )

    paper = Paper(
        subject_id=data.subject_id,
        syllabus_id=data.syllabus_id,
        blueprint_id=data.blueprint_id,
        title=data.title,
        exam_type=data.exam_type,
        duration=data.duration,
        total_marks=data.total_marks,
        total_questions=len(content),
        paper_content=json.dumps(content),
        validation_result=json.dumps(
            {
                "question_count": len(content),
                "status": "generated",
            }
        ),
        status="draft",
        generated_by=user.id,
    )

    db.add(paper)
    db.commit()
    db.refresh(paper)

    return {
        "id": paper.id,
        "content": content,
        "message": "Paper generated",
    }


@router.get("/papers")
def list_papers(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    papers = db.query(Paper).all()

    return [
        {
            "id": p.id,
            "title": p.title,
            "subject_id": p.subject_id,
            "total_marks": p.total_marks,
            "total_questions": p.total_questions,
            "status": p.status,
        }
        for p in papers
    ]


@router.get("/papers/{pid}")
def get_paper(
    pid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    if not paper:
        raise HTTPException(
            404,
            "Paper not found"
        )

    return {
        "id": paper.id,
        "title": paper.title,
        "content": json.loads(
            paper.paper_content or "[]"
        ),
        "total_marks": paper.total_marks,
        "duration": paper.duration,
        "status": paper.status,
    }


@router.get("/papers/{pid}/preview")
def preview_paper(
    pid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    return get_paper(
        pid,
        db,
        user
    )


@router.put("/papers/{pid}")
def update_paper(
    pid: int,
    data: PaperGenIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    if not paper:
        raise HTTPException(
            404,
            "Paper not found"
        )

    for key, value in data.model_dump().items():

        if key != "subject_id" and hasattr(
            paper,
            key
        ):
            setattr(
                paper,
                key,
                value
            )

    db.commit()

    return {
        "message": "Paper updated"
    }


@router.delete("/papers/{pid}")
def delete_paper(
    pid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    if not paper:
        raise HTTPException(
            404,
            "Paper not found"
        )

    db.delete(paper)
    db.commit()

    return {
        "message": "Paper deleted"
    }


@router.post("/papers/{pid}/validate")
def validate_paper(
    pid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    if not paper:
        raise HTTPException(
            404,
            "Paper not found"
        )

    content = json.loads(
        paper.paper_content or "[]"
    )

    result = {
        "question_count": len(content),
        "total_marks": sum(
            x.get("marks", 0)
            for x in content
        ),
        "valid": bool(content),
    }

    paper.validation_result = json.dumps(
        result
    )

    db.commit()

    return result


@router.post("/papers/{pid}/finalize")
def finalize_paper(
    pid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    if not paper:
        raise HTTPException(
            404,
            "Paper not found"
        )

    paper.status = "final"

    vault_item = VaultItem(
        paper_id=paper.id,
        created_by=user.id,
        action="finalized",
    )

    db.add(vault_item)
    db.commit()

    return {
        "message": "Paper finalized"
    }


@router.post("/papers/{pid}/regenerate-question")
def regenerate_question(
    pid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    if not paper:
        raise HTTPException(
            404,
            "Paper not found"
        )

    content = json.loads(
        paper.paper_content or "[]"
    )

    if not content:
        raise HTTPException(
            400,
            "Paper has no questions"
        )

    return {
        "message": (
            "Regeneration is available after "
            "selecting a replacement from the question bank."
        ),
        "content": content,
    }


@router.post("/papers/{pid}/replace-question")
def replace_question(
    pid: int,
    number: int,
    question_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    paper = db.get(
        Paper,
        pid
    )

    question = db.get(
        Question,
        question_id
    )

    if not paper or not question:
        raise HTTPException(
            404,
            "Paper or question not found"
        )

    content = json.loads(
        paper.paper_content or "[]"
    )

    for item in content:

        if item.get("number") == number:

            item.update(
                {
                    "question": question.question_text,
                    "marks": question.marks,
                    "unit": question.unit,
                    "topic": question.topic,
                    "difficulty": question.difficulty,
                    "bloom_level": question.bloom_level,
                    "co": question.course_outcome,
                }
            )

    paper.paper_content = json.dumps(
        content
    )

    db.commit()

    return {
        "message": "Question replaced",
        "content": content,
    }


# =========================================================
# ANSWER KEY
# =========================================================

@router.post("/answer-keys/generate")
def generate_key(
    data: AnswerGenIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    try:
        answer, points = generate_answer(
            data.question_text,
            data.marks
        )
    except Exception:
        from ..services.academic_kb import get_academic_answer
        answer, points = get_academic_answer(data.question_text, data.marks)

    ak_id = None
    try:
        answer_key = AnswerKey(
            paper_id=data.paper_id,
            question_id=data.question_id,
            question_text=data.question_text,
            answer=answer,
            key_points=json.dumps(points),
            marks=data.marks,
            created_by=user.id if user else None,
        )
        db.add(answer_key)
        db.commit()
        db.refresh(answer_key)
        ak_id = answer_key.id
    except Exception:
        db.rollback()

    return {
        "id": ak_id or 1,
        "answer": answer,
        "key_points": points,
    }


@router.get("/answer-keys")
def list_keys(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    answer_keys = (
        db.query(AnswerKey)
        .all()
    )

    return [
        {
            "id": a.id,
            "paper_id": a.paper_id,
            "question_text": a.question_text,
            "marks": a.marks,
            "status": a.status,
        }
        for a in answer_keys
    ]


@router.put("/answer-keys/{aid}")
def update_answer_key(
    aid: int,
    data: AnswerGenIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    answer_key = db.get(
        AnswerKey,
        aid
    )

    if not answer_key:
        raise HTTPException(
            404,
            "Answer key not found"
        )

    answer, points = generate_answer(
        data.question_text,
        data.marks
    )

    answer_key.question_text = (
        data.question_text
    )

    answer_key.marks = data.marks
    answer_key.answer = answer
    answer_key.key_points = json.dumps(points)
    answer_key.faculty_edited = True

    db.commit()

    return {
        "id": answer_key.id,
        "answer": answer,
        "key_points": points,
    }


@router.post("/answer-keys/{aid}/regenerate")
def regenerate_key(
    aid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    answer_key = db.get(
        AnswerKey,
        aid
    )

    if not answer_key:
        raise HTTPException(
            404,
            "Answer key not found"
        )

    answer, points = generate_answer(
        answer_key.question_text,
        answer_key.marks
    )

    answer_key.answer = answer
    answer_key.key_points = json.dumps(points)
    answer_key.ai_generated = True

    db.commit()

    return {
        "id": answer_key.id,
        "answer": answer,
        "key_points": points,
    }


@router.post("/answer-keys/{aid}/validate")
def validate_key(
    aid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    answer_key = db.get(
        AnswerKey,
        aid
    )

    if not answer_key:
        raise HTTPException(
            404,
            "Answer key not found"
        )

    answer_key.validation_status = (
        "valid"
        if answer_key.answer.strip()
        else "invalid"
    )

    db.commit()

    return {
        "valid": (
            answer_key.validation_status
            == "valid"
        ),
        "validation_status":
            answer_key.validation_status,
    }


@router.delete("/answer-keys/{aid}")
def delete_key(
    aid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    answer_key = db.get(
        AnswerKey,
        aid
    )

    if not answer_key:
        raise HTTPException(
            404,
            "Answer key not found"
        )

    db.delete(answer_key)
    db.commit()

    return {
        "message": "Answer key deleted"
    }


@router.post("/answer-keys/{aid}/finalize")
def finalize_key(
    aid: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    answer_key = db.get(
        AnswerKey,
        aid
    )

    if not answer_key:
        raise HTTPException(
            404,
            "Answer key not found"
        )

    answer_key.status = "final"
    answer_key.validation_status = "valid"

    db.commit()

    return {
        "message": "Answer key finalized"
    }


@router.post("/answer-keys/upload")
async def upload_answer_key(
    question_text: str,
    marks: int = 1,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    text = (
        await file.read()
    ).decode(
        "utf-8",
        errors="ignore"
    )

    answer_key = AnswerKey(
        question_text=question_text,
        answer=text,
        marks=marks,
        ai_generated=False,
        created_by=user.id,
    )

    db.add(answer_key)
    db.commit()
    db.refresh(answer_key)

    return {
        "id": answer_key.id,
        "answer": answer_key.answer,
        "message": "Existing answer key uploaded",
    }


# =========================================================
# PREVIOUS PAPER ANALYZER
# =========================================================

@router.post("/previous-papers/analyze")
def analyze_previous(
    subject_id: int,
    title: str = "Previous Paper",
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    analysis = {
        "frequently_asked": [],
        "repeated_topics": [],
        "important_units": [],
        "difficulty_trends": {},
    }

    previous_paper = PreviousPaper(
        subject_id=subject_id,
        title=title,
        analysis=json.dumps(analysis),
        created_by=user.id,
    )

    db.add(previous_paper)
    db.commit()
    db.refresh(previous_paper)

    return {
        "id": previous_paper.id,
        "analysis": analysis,
    }


# =========================================================
# PAPER VAULT
# =========================================================

@router.get("/vault")
def vault(
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    items = (
        db.query(VaultItem)
        .order_by(
            VaultItem.id.desc()
        )
        .all()
    )

    return [
        {
            "id": v.id,
            "paper_id": v.paper_id,
            "version": v.version,
            "action": v.action,
            "notes": v.notes,
            "created_at": v.created_at,
        }
        for v in items
    ]


# =========================================================
# AI EXAM ASSISTANT
# =========================================================

class ChatIn(BaseModel):
    message: str
    subject_id: int | None = None


@router.post("/assistant/chat")
def assistant_chat(
    data: ChatIn,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    question = str(data.message or "").strip()

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Please enter a question."
        )

    context = ""
    if data.subject_id:
        subject = db.get(Subject, data.subject_id)
        if subject:
            context += f"Subject: {subject.name} ({subject.code})\n"
            syllabus = db.query(Syllabus).filter(Syllabus.subject_id == subject.id).first()
            if syllabus:
                context += f"Syllabus Context: {syllabus.extracted_text[:3000]}\n"

    try:
        answer = assistant_answer(question, context=context)
    except Exception:
        from ..services.academic_kb import get_assistant_response
        answer = get_assistant_response(question)

    return {"answer": answer}

@router.put("/questions/{qid}")
def update_question(qid: int, data: dict, db: Session = Depends(get_db), user=Depends(current_user)):
    q = db.get(Question, qid)
    if not q:
        raise HTTPException(404, "Question not found")
    for k, v in data.items():
        if hasattr(q, k) and k != "id":
            setattr(q, k, v)
    db.commit()
    return {"message": "Question updated"}

