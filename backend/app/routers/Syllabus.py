from fastapi import APIRouter, Depends, File, Form, UploadFile, HTTPException
from pydantic import BaseModel
from sqlalchemy.orm import Session
from ..db import get_db
from ..models import Syllabus, Subject

import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

router = APIRouter(
    prefix="/syllabus",
    tags=["Syllabus"]
)


# ============================================================
# REQUEST MODEL FOR MANUAL SYLLABUS
# ============================================================

class ManualSyllabusRequest(BaseModel):
    subject_id: int
    name: str
    description: str = ""
    content: str


# ============================================================
# UNIT DETECTION
# ============================================================

def is_unit_heading(line: str) -> bool:
    line = line.strip()

    if not line:
        return False

    patterns = [
        r"^unit\s*[-:]?\s*[ivxlcdm0-9]+",
        r"^unit\s+[ivxlcdm]+",
        r"^module\s*[-:]?\s*[ivxlcdm0-9]+",
        r"^chapter\s*[-:]?\s*[ivxlcdm0-9]+",
        r"^[ivxlcdm]+\s*[\.\):\-]\s*.+",
        r"^\d+\s*[\.\):\-]\s*.+",
    ]

    return any(
        re.match(pattern, line, re.IGNORECASE)
        for pattern in patterns
    )


def clean_unit_name(line: str) -> str:
    line = line.strip()

    line = re.sub(
        r"^(unit|module|chapter)\s*[-:]?\s*",
        "",
        line,
        flags=re.IGNORECASE
    )

    line = re.sub(
        r"^[ivxlcdm0-9]+\s*[\.\):\-]\s*",
        "",
        line,
        flags=re.IGNORECASE
    )

    return line.strip(" :-")


# ============================================================
# TOPIC EXTRACTION
# ============================================================

def extract_topic(line: str):
    line = line.strip()

    if not line:
        return None

    # Remove bullet symbols
    line = re.sub(
        r"^[•●▪◦\-*]+\s*",
        "",
        line
    )

    # Remove numbered topic format
    # Example:
    # 1. Introduction
    # 1.1 Basics
    # 1.2 Intelligent Agents
    line = re.sub(
        r"^\d+(?:\.\d+)*[\.\)]?\s+",
        "",
        line
    )

    # Remove lettered topics
    # Example:
    # A. Introduction
    # B. Basics
    line = re.sub(
        r"^[A-Za-z][\.\)]\s+",
        "",
        line
    )

    line = line.strip()

    if not line:
        return None

    return line


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(file_bytes: bytes) -> str:
    try:
        import pdfplumber

        text_parts = []

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()

                if text:
                    text_parts.append(text)

        return "\n".join(text_parts)

    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read PDF file: {str(e)}"
        )


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(file_bytes: bytes) -> str:
    try:
        from docx import Document

        document = Document(io.BytesIO(file_bytes))

        text_parts = []

        # Paragraphs
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                text_parts.append(text)

        # Tables
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    text_parts.append(" ".join(row_text))

        return "\n".join(text_parts)

    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read DOCX file: {str(e)}"
        )


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt_text(file_bytes: bytes) -> str:

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    ]

    for encoding in encodings:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise HTTPException(
        status_code=400,
        detail="Unable to read TXT file."
    )


# ============================================================
# FIND LIBREOFFICE
# ============================================================

def find_soffice():

    possible_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]

    for path in possible_paths:
        if os.path.exists(path):
            return path

    return shutil.which("soffice")


# ============================================================
# OLD DOC EXTRACTION
# ============================================================

def extract_doc_text(file_bytes: bytes) -> str:

    soffice = find_soffice()

    if not soffice:
        raise HTTPException(
            status_code=400,
            detail=(
                "DOC files require LibreOffice. "
                "Please install LibreOffice or upload DOCX/PDF/TXT."
            )
        )

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = Path(temp_dir) / "syllabus.doc"

        input_path.write_bytes(file_bytes)

        try:

            result = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    temp_dir,
                    str(input_path),
                ],
                capture_output=True,
                text=True,
                timeout=60,
            )

            if result.returncode != 0:
                raise Exception(result.stderr)

            output_path = Path(temp_dir) / "syllabus.txt"

            if not output_path.exists():
                raise Exception("Converted TXT file was not created.")

            return output_path.read_text(
                encoding="utf-8",
                errors="ignore"
            )

        except Exception as e:

            raise HTTPException(
                status_code=400,
                detail=f"Unable to read DOC file: {str(e)}"
            )


# ============================================================
# FILE TEXT EXTRACTION
# ============================================================

def extract_file_text(filename: str, file_bytes: bytes) -> str:

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    if extension == ".doc":
        return extract_doc_text(file_bytes)

    if extension == ".txt":
        return extract_txt_text(file_bytes)

    raise HTTPException(
        status_code=400,
        detail="Only PDF, DOC, DOCX and TXT files are supported."
    )


# ============================================================
# DETECT UNITS AND TOPICS
# ============================================================

def detect_units(content: str):

    lines = content.splitlines()

    units = []

    current_unit = None

    for raw_line in lines:

        line = raw_line.strip()

        if not line:
            continue

        # Check if line is a unit heading
        if is_unit_heading(line):

            unit_name = clean_unit_name(line)

            current_unit = {
                "unit": unit_name,
                "topics": []
            }

            units.append(current_unit)

            continue

        # Extract topic
        topic = extract_topic(line)

        if not topic:
            continue

        # Ignore common document headings
        ignored = [
            "syllabus",
            "course syllabus",
            "contents",
            "table of contents"
        ]

        if topic.lower() in ignored:
            continue

        if current_unit is None:

            current_unit = {
                "unit": "Syllabus Topics",
                "topics": []
            }

            units.append(current_unit)

        current_unit["topics"].append(topic)

    return units


# ============================================================
# ANALYZE SYLLABUS
# ============================================================

def analyze_content(
    content: str,
    subject_name: str = "",
    description: str = ""
):

    units = detect_units(content)

    topic_count = sum(
        len(unit["topics"])
        for unit in units
    )

    detected_subject = subject_name.strip()

    # Try to detect subject name if not provided
    if not detected_subject:

        patterns = [
            r"subject\s*name\s*[:\-]\s*(.+)",
            r"course\s*name\s*[:\-]\s*(.+)",
            r"subject\s*[:\-]\s*(.+)",
            r"course\s*[:\-]\s*(.+)"
        ]

        for pattern in patterns:

            match = re.search(
                pattern,
                content,
                re.IGNORECASE
            )

            if match:
                detected_subject = match.group(1).strip()
                break

    if not detected_subject:
        detected_subject = "Unknown Subject"

    return {
        "subject_name": detected_subject,
        "description": description.strip(),
        "units_found": len(units),
        "topics_found": topic_count,
        "units": units
    }


# ============================================================
# GET ALL SYLLABI
# ============================================================

@router.get("/")
def get_syllabi(
    db: Session = Depends(get_db)
):

    syllabi = (
        db.query(Syllabus)
        .order_by(Syllabus.id.desc())
        .all()
    )

    result = []

    for syllabus in syllabi:

        result.append({
            "id": syllabus.id,
            "subject_id": syllabus.subject_id,
            "name": syllabus.name,
            "description": syllabus.description,
            "status": syllabus.status,
            "created_at": syllabus.created_at,
        })

    return result


# ============================================================
# MANUAL SYLLABUS ENTRY
# ============================================================

@router.post("/")
def create_manual_syllabus(
    data: ManualSyllabusRequest,
    db: Session = Depends(get_db)
):

    # Check subject
    subject = (
        db.query(Subject)
        .filter(Subject.id == data.subject_id)
        .first()
    )

    if not subject:

        raise HTTPException(
            status_code=404,
            detail="Subject not found."
        )

    # Validate syllabus name
    if not data.name.strip():

        raise HTTPException(
            status_code=400,
            detail="Syllabus name is required."
        )

    # Validate content
    if not data.content.strip():

        raise HTTPException(
            status_code=400,
            detail="Syllabus content is required."
        )

    # Detect units
    units = detect_units(data.content)

    # Save to database
    syllabus = Syllabus(
        subject_id=data.subject_id,
        name=data.name.strip(),
        description=data.description.strip(),
        content=data.content.strip(),
        units=json.dumps(
            units,
            ensure_ascii=False
        ),
        status="Ready"
    )

    db.add(syllabus)

    db.commit()

    db.refresh(syllabus)

    return {
        "message": "Syllabus saved successfully.",
        "syllabus": {
            "id": syllabus.id,
            "subject_id": syllabus.subject_id,
            "name": syllabus.name,
            "description": syllabus.description,
            "status": syllabus.status
        }
    }


# ============================================================
# UPLOAD SYLLABUS
# ============================================================

@router.post("/upload")
async def upload_syllabus(
    file: UploadFile = File(...),
    subject_id: int = Form(...),
    name: str = Form(...),
    description: str = Form(""),
    db: Session = Depends(get_db)
):

    # Check subject
    subject = (
        db.query(Subject)
        .filter(Subject.id == subject_id)
        .first()
    )

    if not subject:

        raise HTTPException(
            status_code=404,
            detail="Subject not found."
        )

    # Validate name
    if not name.strip():

        raise HTTPException(
            status_code=400,
            detail="Syllabus name is required."
        )

    # Read file
    file_bytes = await file.read()

    if not file_bytes:

        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty."
        )

    # Extract text
    content = extract_file_text(
        file.filename,
        file_bytes
    )

    if not content.strip():

        raise HTTPException(
            status_code=400,
            detail="No readable text found in the uploaded file."
        )

    # Detect units
    units = detect_units(content)

    # Save to database
    syllabus = Syllabus(
        subject_id=subject_id,
        name=name.strip(),
        description=description.strip(),
        content=content.strip(),
        units=json.dumps(
            units,
            ensure_ascii=False
        ),
        status="Ready"
    )

    db.add(syllabus)

    db.commit()

    db.refresh(syllabus)

    return {
        "message": "Syllabus uploaded successfully.",
        "syllabus": {
            "id": syllabus.id,
            "subject_id": syllabus.subject_id,
            "name": syllabus.name,
            "description": syllabus.description,
            "status": syllabus.status
        }
    }


# ============================================================
# ANALYZE SAVED SYLLABUS
# ============================================================

@router.post("/{syllabus_id}/analyze")
def analyze_syllabus(
    syllabus_id: int,
    db: Session = Depends(get_db)
):

    syllabus = (
        db.query(Syllabus)
        .filter(Syllabus.id == syllabus_id)
        .first()
    )

    if not syllabus:

        raise HTTPException(
            status_code=404,
            detail="Syllabus not found."
        )

    # Get subject
    subject = None

    if syllabus.subject_id:

        subject = (
            db.query(Subject)
            .filter(Subject.id == syllabus.subject_id)
            .first()
        )

    subject_name = ""

    if subject:
        subject_name = subject.name

    # Analyze content
    analysis = analyze_content(
        syllabus.content or "",
        subject_name,
        syllabus.description or ""
    )

    # Update stored units
    syllabus.units = json.dumps(
        analysis["units"],
        ensure_ascii=False
    )

    syllabus.status = "Analyzed"

    db.commit()

    return {
        "message": "Syllabus analyzed successfully.",
        "analysis": analysis
    }